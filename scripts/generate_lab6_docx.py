#!/usr/bin/env python3
import os
import shutil
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def create_lab6_doc():
    template_path = "docs/evidence/TV1/Lab05_2312741_NguyenThanhTam.docx"
    doc = docx.Document(template_path)

    # 1. Update header paragraphs
    for p in doc.paragraphs:
        if "Lab: 05" in p.text or "Lab: 04" in p.text or "Lab:" in p.text:
            p.text = "Lab: 06 (Tuần 6)\tTừ ngày: 08/10/2026\tđến ngày: 14/10/2026"
            for r in p.runs:
                r.bold = True

    tasks = [
        {
            "stt": "1",
            "title": "[Kiểm thử hồi quy Backend] Chạy toàn diện 446 bài test tự động backend bảo đảm 100% Green",
            "done": [
                "- Kiểm thử bộ test chính tests/CulinaryBlog.Tests: 441/441 bài test tự động vượt qua 100% trong thời gian 2m55s.",
                "- Kiểm thử bộ test tương tranh tests/concurrency-spike: 5/5 bài test tương tranh vượt qua 100%.",
                "- Tổng số bài test backend đạt 446/446 tests PASS 100% (0 Failed, 0 Skipped).",
                "- Duy trì Code Coverage tầng CulinaryBlog.Application đạt 96.17% (vượt xa chỉ tiêu 80%)."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/tests/CulinaryBlog.Tests",
            "progress": "100%"
        },
        {
            "stt": "2",
            "title": "[Kiểm thử hồi quy Frontend] Chạy 86 bài test Jest và biên dịch Next.js 15 App Router production sạch sẽ",
            "done": [
                "- Chạy Jest test suite trong src/frontend: Đạt 15/15 test suites passed, 86/86 tests passed 100%.",
                "- Biên dịch npm run build cho Next.js 15.1.11 thành công toàn bộ 16/16 routes tĩnh và động mà không gặp lỗi typecheck hay compilation nào."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend",
            "progress": "100%"
        },
        {
            "stt": "3",
            "title": "[Chuẩn hóa mã nguồn & Rà soát bảo mật] Xác thực định dạng code và kiểm tra rò rỉ secret",
            "done": [
                "- Chạy dotnet format CulinaryBlog.sln --verify-no-changes --no-restore: Exit code 0, không có vi phạm formatting nào trong toàn bộ solution.",
                "- Chạy bash deploy/scan-secrets.sh: Kết quả OK: không phát hiện secret bị gửi vào repo.",
                "- Chạy bash deploy/scan-secrets.test.sh: Đạt 9/9 pass, 0 fail."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/deploy/scan-secrets.sh",
            "progress": "100%"
        },
        {
            "stt": "4",
            "title": "[Nghiệm thu 24/24 Kỹ năng] Rà soát và đánh giá ma trận K01–K24 cho toàn bộ 4 thành viên",
            "done": [
                "- Rà soát mã nguồn, cấu hình và bài test của TV1, TV2, TV3, TV4 đối chiếu 24 nhóm kỹ năng.",
                "- Xác nhận cả 4 thành viên đạt 24/24 nhóm kỹ năng có code/config và test kiểm chứng thực tế.",
                "- Review và merge các PR chốt Tuần 6: PR #34, PR #35 (TV3), PR #36 (TV4), PR #31 (TV2) vào nhánh chính main."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/docs/PHAN_CHIA_CONG_VIEC_6_TUAN.md",
            "progress": "100%"
        },
        {
            "stt": "5",
            "title": "[Kịch bản Demo phân hệ Tài khoản & Hồ sơ] Xây dựng và diễn tập kịch bản demo 15–20 phút",
            "done": [
                "- Thiết lập kịch bản 5 phần hoàn chỉnh: (1) Kiến trúc & cấu hình PBKDF2 HMAC-SHA512 + JWT HS512; (2) Luồng đăng ký, đăng nhập & chống user enumeration; (3) Dashboard hồ sơ, đổi mật khẩu & thu hồi token đa phiên; (4) Chống brute-force lockout (5 lần) + Rate limit + XSS; (5) Diễn tập khôi phục sự cố dữ liệu BCP/DR SHA-256 (1.8s).",
                "- Diễn tập chạy thử nghiệm thành công 100% trên môi trường local và staging."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/docs/evidence/TV1/TUAN_6.md",
            "progress": "100%"
        },
        {
            "stt": "6",
            "title": "[Bàn giao đồ án Release v1.0.0] Đóng gói hồ sơ tài liệu và cập nhật README.md",
            "done": [
                "- Cập nhật toàn diện README.md với tiến độ nghiệm thu Cổng G7, bảng số liệu kiểm thử chính xác, hướng dẫn chạy app và hướng dẫn chạy demo cho Giảng viên.",
                "- Cập nhật tiến độ 4 thành viên trong docs/PHAN_CHIA_CONG_VIEC_6_TUAN.md.",
                "- Hoàn tất bộ tài liệu dự án gồm SRS v1.1.1, ERD, 6 Hợp đồng kỹ thuật, 8 ADRs, 2 Runbooks vận hành."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/README.md",
            "progress": "100%"
        }
    ]

    table = doc.tables[0]
    # Keep header row (index 0), remove all other rows
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[1]._tr)

    for task in tasks:
        row = table.add_row()

        # Col 0: STT
        c0 = row.cells[0]
        c0.width = Inches(0.485)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        r0 = p0.add_run(task["stt"])
        r0.bold = True
        r0.font.size = Pt(10.0)

        # Col 1: Công việc được giao
        c1 = row.cells[1]
        c1.width = Inches(4.159)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(4.0)
        r_title = p1.add_run(task["title"])
        r_title.bold = True
        r_title.font.size = Pt(9.5)

        # Đã hoàn thành header
        p_done_head = c1.add_paragraph()
        p_done_head.paragraph_format.space_after = Pt(2.0)
        r_dh = p_done_head.add_run("Đã hoàn thành:")
        r_dh.bold = True
        r_dh.font.size = Pt(9.0)

        for d_item in task["done"]:
            p_d = c1.add_paragraph()
            p_d.paragraph_format.space_after = Pt(2.0)
            r_d = p_d.add_run(d_item)
            r_d.font.size = Pt(8.5)

        # Chưa hoàn thành header
        p_not_head = c1.add_paragraph()
        p_not_head.paragraph_format.space_before = Pt(3.0)
        p_not_head.paragraph_format.space_after = Pt(2.0)
        r_nh = p_not_head.add_run("Chưa hoàn thành:")
        r_nh.bold = True
        r_nh.font.size = Pt(9.0)

        for nd_item in task["not_done"]:
            p_nd = c1.add_paragraph()
            p_nd.paragraph_format.space_after = Pt(2.0)
            r_nd = p_nd.add_run(nd_item)
            r_nd.font.size = Pt(8.5)

        # Col 2: Link
        c2 = row.cells[2]
        c2.width = Inches(4.216)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(task["link"])
        r2.font.size = Pt(7.5)

        # Col 3: Tiến độ %
        c3 = row.cells[3]
        c3.width = Inches(1.0)
        p3 = c3.paragraphs[0]
        p3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        r3 = p3.add_run(task["progress"])
        r3.bold = True
        r3.font.size = Pt(10.0)

    out_paths = [
        "docs/evidence/TV1/Lab6_2312741_NguyenThanhTam.docx",
        "docs/evidence/TV1/Lab06_2312741_NguyenThanhTam.docx",
        os.path.expanduser("~/Downloads/Lab6_2312741_NguyenThanhTam.docx"),
        os.path.expanduser("~/Downloads/Lab06_2312741_NguyenThanhTam.docx")
    ]

    for p in out_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        doc.save(p)
        print(f"Generated: {p} (size: {os.path.getsize(p)} bytes)")

if __name__ == "__main__":
    create_lab6_doc()
