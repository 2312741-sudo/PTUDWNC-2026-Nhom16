#!/usr/bin/env python3
import os
import shutil
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def create_lab6_doc_tv2():
    template_path = "docs/evidence/TV2/Lab04_2312796_NgoQuocTruongVi.docx"
    doc = docx.Document(template_path)

    # 1. Update header paragraphs
    for p in doc.paragraphs:
        if "Lab: 04" in p.text or "Lab: 4" in p.text or "Lab: 05" in p.text:
            p.text = "Lab: 06 (Tuần 6)\tTừ ngày: 07/10/2026\tđến ngày: 14/10/2026"
            for r in p.runs:
                r.bold = True
        elif "Lab 04" in p.text or "Lab 05" in p.text:
            p.text = p.text.replace("Lab 04", "Lab 06").replace("Lab 05", "Lab 06").replace("Tuần 4", "Tuần 6").replace("Tuần 5", "Tuần 6")

    tasks = [
        {
            "stt": "1",
            "title": "[Yêu cầu 1 Lab 6 — Lớp tiện ích xử lý lỗi & gọi API] Cài đặt hệ thống phân cấp lỗi chuẩn RFC 7807/9457 và HTTP Client",
            "done": [
                "- Cài đặt hệ thống phân loại lỗi chuyên sâu (src/frontend/src/lib/errors.ts) bao gồm lớp gốc AppError, lớp ApiError (lưu trữ mã trạng thái, tiêu đề, chi tiết, mã lỗi hệ thống), lớp ValidationError xử lý lỗi chi tiết từng trường form (getFieldError), cùng các lớp lỗi HTTP đặc thù: UnauthorizedError (401), ForbiddenError (403), NotFoundError (404), NetworkError và TimeoutError.",
                "- Xây dựng tiện ích HttpClient (src/frontend/src/lib/httpClient.ts) hỗ trợ đầy đủ các phương thức HTTP (get, post, put, patch, delete), tự động phân giải Base URL SSR/Client, tích hợp Request Interceptor tự động nạp Bearer Token từ auth-store, Response Interceptor tự động bóc tách dữ liệu và ném lỗi có cấu trúc, kèm bộ điều khiển ngắt kết nối khi quá thời gian chờ (Timeout Controller 15s).",
                "- Viết 13 bài test tự động cho errors.test.ts và httpClient.test.ts đạt kết quả Green 100%."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/2312796_NgoQuocTruongVi_lab6/src/frontend/src/lib/httpClient.ts",
            "progress": "100%"
        },
        {
            "stt": "2",
            "title": "[Yêu cầu 2 Lab 6 — Cài đặt gói thư viện shadcn/ui] Cấu hình shadcn/ui và xây dựng bộ giao diện thành phần",
            "done": [
                "- Khởi tạo tệp cấu hình chuẩn components.json định nghĩa phong cách default, màu nền tảng emerald, hỗ trợ React Server Components (RSC) và path aliases chuẩn mực @/components, @/lib/utils, @/components/ui.",
                "- Xây dựng hàm hợp nhất class cn (src/frontend/src/lib/utils.ts) kết hợp clsx và tailwind-merge loại trừ xung đột class CSS Tailwind, kèm các hàm định dạng tiền tệ VND (formatVND) và thời gian chuẩn Việt Nam (formatDate).",
                "- Xây dựng hoàn chỉnh bộ UI components theo chuẩn shadcn/ui trong src/frontend/src/components/ui/:",
                "  + button.tsx: Hỗ trợ biến thể màu default, emerald, destructive, outline, secondary, ghost, link; kích cỡ sm, default, lg, icon; hỗ trợ cờ isLoading spinner và thuộc tính asChild liên kết linh hoạt với next/link.",
                "  + input.tsx: Ô nhập dữ liệu tích hợp icon và thông báo lỗi validation trực quan.",
                "  + card.tsx: Bộ thẻ cấu trúc phân tầng (Card, CardHeader, CardTitle, CardDescription, CardContent, CardFooter).",
                "  + badge.tsx: Nhãn trạng thái sắc nét với các biến thể success, warning, info, destructive.",
                "  + dialog.tsx: Hộp thoại Modal có hiệu ứng mờ nền (backdrop blur) và khóa cuộn trang an toàn.",
                "  + skeleton.tsx: Khung xương tải trang chuyển động mềm mại animate-pulse.",
                "  + alert.tsx: Hộp thông báo ngữ nghĩa kèm icon cảnh báo tương ứng."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/2312796_NgoQuocTruongVi_lab6/src/frontend/src/components/ui/button.tsx",
            "progress": "100%"
        },
        {
            "stt": "3",
            "title": "[Yêu cầu 3 Lab 6 — Store quản lý trạng thái auth-store.ts] Quản lý trạng thái ứng dụng với Zustand và LocalStorage Sync",
            "done": [
                "- Cài đặt và tích hợp thư viện quản lý trạng thái hiện đại zustand vào dự án frontend.",
                "- Xây dựng kho lưu trữ trạng thái đăng nhập useAuthStore (src/frontend/src/store/auth-store.ts) quản lý thông tin tài khoản người dùng (User), JWT access token, refresh token, trạng thái xác thực isAuthenticated, trạng thái khởi tạo isInitialized và cờ đang xử lý isLoading.",
                "- Triển khai cơ chế đồng bộ hai chiều với localStorage, phát sự kiện tùy biến auth-state-change giúp mọi component nhận diện thay đổi tức thì mà không cần tải lại trang.",
                "- Cài đặt các hàm hành động: setAuth, setAuthFromResponse, updateUser, logout, initialize và hàm kiểm tra phân quyền hasRole(user, role).",
                "- Bổ sung useUIStore (src/frontend/src/store/ui-store.ts) quản lý thông báo Toast tự động biến mất và trạng thái đóng/mở menu di động.",
                "- Tích hợp useAuthStore vào thanh điều hướng Header.tsx, trang đăng nhập app/auth/login/page.tsx và trang đăng ký app/auth/register/page.tsx.",
                "- Viết 6 bài unit test trong auth-store.test.ts pass 100%."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/2312796_NgoQuocTruongVi_lab6/src/frontend/src/store/auth-store.ts",
            "progress": "100%"
        },
        {
            "stt": "4",
            "title": "[Yêu cầu 4 Lab 6 — Thiết kế giao diện (tĩnh) cho blog app] Xây dựng giao diện Blog App cao cấp và Mock Data Fallback",
            "done": [
                "- Xây dựng tệp dữ liệu tĩnh hoàn chỉnh (src/frontend/src/lib/mockData.ts) bao gồm 6 danh mục ẩm thực phong phú (MOCK_CATEGORIES), 6 bài viết công thức chuẩn vị (MOCK_RECIPES), chi tiết công thức nấu ăn chuyên sâu (MOCK_RECIPE_DETAIL) và thuật toán phân trang tĩnh getMockPagedRecipes.",
                "- Nâng cấp toàn diện giao diện Trang chủ (src/frontend/src/app/page.tsx):",
                "  + Hero Section ấn tượng với hiệu ứng radial gradient, nhãn nổi bật ẩm thực Việt Nam và tiêu đề lớn truyền cảm hứng.",
                "  + Thanh tìm kiếm nhanh kèm các thẻ gợi ý ('phở bò', 'cá kho tộ', 'bún chả', 'bánh ngọt').",
                "  + Lưới trình bày danh mục ẩm thực (CategoryCard) và các công thức nấu ăn mới nhất (RecipeCard).",
                "  + Dải thông số thống kê ấn tượng (500+ món ngon, 25+ danh mục, 10,000+ độc giả, 4.9 sao).",
                "  + Khối tính năng cốt lõi giới thiệu công nghệ tìm kiếm FTS không dấu, công thức chi tiết từng bước và bảo mật.",
                "  + Form đăng ký nhận bản tin ẩm thực tuần qua email (Newsletter Subscription).",
                "- Tích hợp cơ chế tự động chuyển sang dữ liệu mẫu tĩnh trong src/frontend/src/lib/api.ts khi Backend API offline, đảm bảo website luôn hiển thị chỉn chu và không bao giờ bị lỗi giao diện trống."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/2312796_NgoQuocTruongVi_lab6/src/frontend/src/app/page.tsx",
            "progress": "100%"
        },
        {
            "stt": "5",
            "title": "[Yêu cầu 5 Lab 6 — Xây dựng các component dùng chung] Hoàn thiện Header, Footer, Pagination, Loading và Error",
            "done": [
                "- Header.tsx: Thanh điều hướng cố định trên cùng, kết nối trực tiếp với useAuthStore để phản hồi trạng thái đăng nhập tức thì, tích hợp thanh tìm kiếm nhanh, nút viết công thức và menu quản trị.",
                "- Footer.tsx: Chân trang đa cột đầy đủ thông tin bản quyền, danh mục, liên kết nhanh, thông điệp bảo trợ ẩm thực và mạng xã hội.",
                "- Pagination.tsx (src/frontend/src/components/Pagination.tsx): Component phân trang dùng chung cao cấp, hỗ trợ cả 2 chế độ tạo URL SSR (createPageUrl) và callback Client State (onPageChange), tích hợp nút Đầu/Trước/Số trang có dấu chấm lửng/Sau/Cuối, tuân thủ tiêu chuẩn trợ năng WCAG 2.1 AA (aria-label, aria-current='page').",
                "- Loading.tsx (src/frontend/src/components/Loading.tsx): Cung cấp LoadingSpinner nhiều kích cỡ, LoadingPage toàn màn hình và bộ khung xương RecipeCardSkeleton, RecipeListSkeleton.",
                "- ErrorMessage.tsx (src/frontend/src/components/ErrorMessage.tsx): Cung cấp ErrorMessage toàn trang có nút Thử lại (Retry) và Về trang chủ, cùng InlineErrorAlert dùng cho các thông báo lỗi tức thời.",
                "- Viết 13 bài test trong CommonComponents.test.tsx kiểm thử hành vi toàn bộ các component dùng chung đạt 100% PASS."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/2312796_NgoQuocTruongVi_lab6/src/frontend/src/components/Pagination.tsx",
            "progress": "100%"
        },
        {
            "stt": "6",
            "title": "[Yêu cầu 6 Lab 6 — Bố cục (Layout) chung cho blog app] Xây dựng RootLayout và các Route Boundaries của Next.js 15",
            "done": [
                "- Cài đặt bố cục gốc toàn diện src/frontend/src/app/layout.tsx với cấu trúc ngữ nghĩa HTML5 (html, body className='antialiased flex flex-col min-h-screen', Header, main className='flex-1', Footer).",
                "- Cấu hình Metadata SEO động, OpenGraph, từ khóa và biểu tượng favicon chuẩn mực.",
                "- Xây dựng trọn bộ các trang biên ranh giới chuyên dụng chuẩn App Router của Next.js 15:",
                "  + src/frontend/src/app/loading.tsx: Tự động kết xuất giao diện tải trang mượt mà khi người dùng điều hướng route.",
                "  + src/frontend/src/app/error.tsx: Bắt lỗi biên ranh giới (Error Boundary) toàn ứng dụng, hiển thị hộp thoại khắc phục và nút khôi phục (Reset).",
                "  + src/frontend/src/app/not-found.tsx: Trang 404 thân thiện thông báo khi bài viết hoặc công thức ẩm thực không tìm thấy."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/2312796_NgoQuocTruongVi_lab6/src/frontend/src/app/layout.tsx",
            "progress": "100%"
        },
        {
            "stt": "7",
            "title": "[Đảm bảo chất lượng toàn diện & Bộ kiểm thử tự động Frontend] 19 Test Suites Green 100% và Biên dịch TypeScript sạch",
            "done": [
                "- Toàn bộ 19 test suites frontend (118 bài test Jest) đạt kết quả PASS 100%.",
                "- Kiểm tra kiểu dữ liệu TypeScript (npx tsc --noEmit) đạt 0 lỗi biên dịch.",
                "- Kiểm tra tương thích Next.js 15 App Router (npm run build) đạt kết quả xuất sắc.",
                "- Hoàn thành tệp báo cáo TUAN_6.md, BAO_CAO_LAB_06.md và mã nguồn sinh tệp Word scripts/generate_lab6_docx_tv2.py."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/2312796_NgoQuocTruongVi_lab6/docs/evidence/TV2/TUAN_6.md",
            "progress": "100%"
        }
    ]

    table = doc.tables[0]

    # Delete existing data rows (keep header row 0)
    while len(table.rows) > 1:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)

    # Add each task as a row
    for task in tasks:
        row = table.add_row()
        # Col 0: STT
        c0 = row.cells[0]
        c0.width = Inches(0.684)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        r0 = p0.add_run(task["stt"])
        r0.bold = True
        r0.font.size = Pt(10.0)

        # Col 1: Công việc được giao
        c1 = row.cells[1]
        c1.width = Inches(4.159)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(4.0)
        r_title = p1.add_run(task["title"])
        r_title.bold = True
        r_title.font.size = Pt(9.5)

        # Đã hoàn thành header
        p_done_head = c1.add_paragraph()
        p_done_head.paragraph_format.space_after = Pt(2.0)
        r_dh = p_done_head.add_run("Đã hoàn thành:")
        r_dh.bold = True
        r_dh.font.size = Pt(9.0)

        for d_item in task["done"]:
            p_d = c1.add_paragraph()
            p_d.paragraph_format.space_after = Pt(2.0)
            r_d = p_d.add_run(d_item)
            r_d.font.size = Pt(8.5)

        # Chưa hoàn thành header
        p_not_head = c1.add_paragraph()
        p_not_head.paragraph_format.space_before = Pt(3.0)
        p_not_head.paragraph_format.space_after = Pt(2.0)
        r_nh = p_not_head.add_run("Chưa hoàn thành:")
        r_nh.bold = True
        r_nh.font.size = Pt(9.0)

        for nd_item in task["not_done"]:
            p_nd = c1.add_paragraph()
            p_nd.paragraph_format.space_after = Pt(2.0)
            r_nd = p_nd.add_run(nd_item)
            r_nd.font.size = Pt(8.5)

        # Col 2: Link
        c2 = row.cells[2]
        c2.width = Inches(4.216)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(task["link"])
        r2.font.size = Pt(7.5)

        # Col 3: Tiến độ %
        c3 = row.cells[3]
        c3.width = Inches(1.0)
        p3 = c3.paragraphs[0]
        p3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        r3 = p3.add_run(task["progress"])
        r3.bold = True
        r3.font.size = Pt(10.0)

    # Save to required destinations
    out_paths = [
        "docs/evidence/TV2/Lab6_2312796_NgoQuocTruongVi.docx",
        "docs/evidence/TV2/Lab06_2312796_NgoQuocTruongVi.docx",
        os.path.expanduser("~/Downloads/Lab6_2312796_NgoQuocTruongVi.docx"),
        os.path.expanduser("~/Downloads/Lab06_2312796_NgoQuocTruongVi.docx")
    ]

    for p in out_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        doc.save(p)
        print(f"Generated: {p} (size: {os.path.getsize(p)} bytes)")

if __name__ == "__main__":
    create_lab6_doc_tv2()
