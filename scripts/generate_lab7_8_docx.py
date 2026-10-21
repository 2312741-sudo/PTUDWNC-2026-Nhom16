#!/usr/bin/env python3
import os
import shutil
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def create_lab_doc(lab_num_str, date_range_str, tasks, out_filenames):
    template_path = "docs/evidence/TV1/Lab05_2312741_NguyenThanhTam.docx"
    doc = docx.Document(template_path)

    # 1. Update header paragraphs
    for p in doc.paragraphs:
        if "Lab: 05" in p.text or "Lab: 04" in p.text or "Lab:" in p.text:
            p.text = f"Lab: {lab_num_str}\t{date_range_str}"
            for r in p.runs:
                r.bold = True

    table = doc.tables[0]
    # Keep header row (index 0), remove all other rows
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[1]._tr)

    for task in tasks:
        row = table.add_row()

        # Col 0: STT
        c0 = row.cells[0]
        c0.width = Inches(0.485)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        r0 = p0.add_run(task["stt"])
        r0.bold = True
        r0.font.size = Pt(10.0)

        # Col 1: Công việc được giao
        c1 = row.cells[1]
        c1.width = Inches(4.159)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(4.0)
        r_title = p1.add_run(task["title"])
        r_title.bold = True
        r_title.font.size = Pt(9.5)

        # Đã hoàn thành header
        p_done_head = c1.add_paragraph()
        p_done_head.paragraph_format.space_after = Pt(2.0)
        r_dh = p_done_head.add_run("Đã hoàn thành:")
        r_dh.bold = True
        r_dh.font.size = Pt(9.0)

        for d_item in task["done"]:
            p_d = c1.add_paragraph()
            p_d.paragraph_format.space_after = Pt(2.0)
            r_d = p_d.add_run(d_item)
            r_d.font.size = Pt(8.5)

        # Chưa hoàn thành header
        p_not_head = c1.add_paragraph()
        p_not_head.paragraph_format.space_before = Pt(3.0)
        p_not_head.paragraph_format.space_after = Pt(2.0)
        r_nh = p_not_head.add_run("Chưa hoàn thành:")
        r_nh.bold = True
        r_nh.font.size = Pt(9.0)

        for nd_item in task["not_done"]:
            p_nd = c1.add_paragraph()
            p_nd.paragraph_format.space_after = Pt(2.0)
            r_nd = p_nd.add_run(nd_item)
            r_nd.font.size = Pt(8.5)

        # Col 2: Link
        c2 = row.cells[2]
        c2.width = Inches(4.216)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(task["link"])
        r2.font.size = Pt(7.5)

        # Col 3: Tiến độ %
        c3 = row.cells[3]
        c3.width = Inches(1.0)
        p3 = c3.paragraphs[0]
        p3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        r3 = p3.add_run(task["progress"])
        r3.bold = True
        r3.font.size = Pt(10.0)

    for p in out_filenames:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        doc.save(p)
        print(f"Generated: {p} (size: {os.path.getsize(p)} bytes)")

def main():
    lab7_tasks = [
        {
            "stt": "1",
            "title": "[Phân hệ Xác thực] Hoàn thành việc cài đặt các tính năng: đăng ký, đăng nhập, đăng xuất, đổi mật khẩu, quên mật khẩu",
            "done": [
                "- Đăng ký tài khoản (POST /api/v1/auth/register): Trang src/app/auth/register/page.tsx, mã hóa PBKDF2 HMAC-SHA512 (100.000 iterations), tự động gán role Author.",
                "- Đăng nhập (POST /api/v1/auth/login): Trang src/app/auth/login/page.tsx, hỗ trợ Google OAuth PKCE, chống brute-force tự động khóa tài khoản sau 5 lần sai (Lockout 15p, HTTP 423), cấp phát cặp JWT HS512 + Refresh Token.",
                "- Đăng xuất (POST /api/v1/auth/logout): Thu hồi Refresh Token trên DB, làm sạch client state tại Header.tsx.",
                "- Đổi mật khẩu (POST /api/v1/auth/change-password): Trang src/app/dashboard/profile/page.tsx, hỗ trợ gửi mã OTP qua email, kích hoạt Token Family Revocation thu hồi mọi token cũ.",
                "- Quên mật khẩu (POST /api/v1/auth/forgot-password): Trang src/app/auth/forgot-password/page.tsx, sinh mật khẩu an toàn gửi về email."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend/src/app/auth",
            "progress": "100%"
        },
        {
            "stt": "2",
            "title": "[Quản lý Hồ sơ] Hoàn thành việc xây dựng trang hồ sơ (profile), cho phép cập nhật hồ sơ",
            "done": [
                "- Xây dựng trang hồ sơ cá nhân tại src/app/dashboard/profile/page.tsx.",
                "- Hiển thị thông tin người dùng từ GET /api/v1/auth/me: Họ và tên, email, vai trò (Role badge: Admin, Author), ảnh đại diện, tiểu sử Bio, ngày tham gia.",
                "- Cập nhật hồ sơ qua PATCH /api/v1/auth/me: Cập nhật Họ và tên, tiểu sử Bio, avatar với xác thực FluentValidation chống XSS.",
                "- Đổi mật khẩu trực tiếp trong trang hồ sơ với đếm ngược OTP cooldown."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend/src/app/dashboard/profile/page.tsx",
            "progress": "100%"
        },
        {
            "stt": "3",
            "title": "[Trang chủ đa khối] Hoàn thành việc xây dựng trang chủ (có đủ header, footer, sidebar hiển thị danh sách category, main content hiển thị danh sách công thức mới, ...)",
            "done": [
                "- Có đầy đủ Header và Footer dùng chung toàn ứng dụng tại layout.tsx, Header.tsx và Footer.tsx.",
                "- Hero Section: Khẩu hiệu truyền cảm hứng, giới thiệu nền tảng và các nút CTA hành động nhanh.",
                "- Khối Sidebar danh mục: Cột bên trái hiển thị danh mục món ăn (Category list), icon chuyên mục, hiệu ứng hover và liên kết nhanh.",
                "- Khối Main Content công thức mới: Cột chính hiển thị các món ăn mới nhất (Newest Recipes Grid qua getRecipes) với RecipeCard trực quan.",
                "- Khối Danh mục nổi bật và Khối Điểm nhấn tính năng (FTS không dấu, Chi tiết từng bước, Kiến trúc tin cậy)."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend/src/app/page.tsx",
            "progress": "100%"
        },
        {
            "stt": "4",
            "title": "[Chi tiết công thức] Hoàn thành việc xây dựng trang xem chi tiết công thức (món ăn)",
            "done": [
                "- Xây dựng trang chi tiết tại src/app/recipes/[slug]/page.tsx theo cơ chế ISR 5 phút và generateStaticParams.",
                "- Hiển thị đầy đủ: Tên món ăn, mô tả, ảnh đại diện có fallback, thời gian nấu, khẩu phần ăn, độ khó, tác giả, danh mục.",
                "- Danh sách nguyên liệu chi tiết (Ingredients List) kèm định lượng và ghi chú.",
                "- Hướng dẫn từng bước (Cooking Steps) có ảnh công đoạn và bộ hẹn giờ đếm ngược (Step Timer).",
                "- Bảng dinh dưỡng (Nutrition Facts) và Schema.org/Recipe JSON-LD chuẩn SEO."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend/src/app/recipes/%5Bslug%5D/page.tsx",
            "progress": "100%"
        },
        {
            "stt": "5",
            "title": "[Tìm kiếm món ăn] Hoàn thành việc cài đặt chức năng tìm kiếm",
            "done": [
                "- Backend: PostgreSQL Full-Text Search qua unaccent, pg_trgm, chỉ mục GIN và thuật toán xếp hạng ts_rank, hỗ trợ tìm kiếm không dấu tối ưu.",
                "- Frontend: Trang tìm kiếm tại src/app/search/page.tsx và thanh tìm kiếm nhanh Header.tsx.",
                "- Bộ lọc đa tiêu chí: Lọc kết hợp từ khóa q, Danh mục, Độ khó, sắp xếp theo ngày/thời gian nấu, phân trang mượt mà.",
                "- Phân lập bảo mật: Tuyệt đối không trả về bài viết Draft trong kết quả tìm kiếm."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend/src/app/search/page.tsx",
            "progress": "100%"
        },
        {
            "stt": "6",
            "title": "[Danh sách Category] Hoàn thành việc xây dựng trang xem danh sách tất cả category",
            "done": [
                "- Trang tổng hợp toàn bộ danh mục tại src/app/categories/page.tsx (ISR 1 giờ, lưới CategoryCard kèm số lượng bài viết).",
                "- Trang chi tiết từng danh mục tại src/app/categories/[slug]/page.tsx kèm phân trang và breadcrumb.",
                "- Trang quản trị danh mục dành cho Admin tại src/app/dashboard/categories/page.tsx."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend/src/app/categories/page.tsx",
            "progress": "100%"
        }
    ]

    lab8_tasks = [
        {
            "stt": "1",
            "title": "[Tải lên & Xử lý ảnh] Hoàn thành việc cài đặt tính năng upload ảnh",
            "done": [
                "- Backend: POST /api/v1/recipes/{id}/images, kiểm tra MIME và Magic Bytes nhị phân (JPEG, PNG, WEBP, GIF), giới hạn <= 5 MiB.",
                "- Lưu trữ RustFS/MinIO/S3, proxy ảnh có xác thực bảo vệ ảnh thuộc bài viết Draft.",
                "- Hangfire Background Job tự động resize ảnh bất đồng bộ thành 2 kích thước chuẩn 300x300 và 800x600 có retry.",
                "- Frontend: Component tải ảnh ImagesStep.tsx với thanh tiến độ upload, xem trước ảnh, đặt ảnh chính Primary và xóa ảnh."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend/src/app/dashboard/recipes/_wizard/ImagesStep.tsx",
            "progress": "100%"
        },
        {
            "stt": "2",
            "title": "[Soạn thảo & Vòng đời Công thức] Hoàn thành việc xây dựng trang thêm mới/cập nhật công thức, thay đổi trạng thái công thức (Draft -> Publish -> Archived/Deleted)",
            "done": [
                "- Trang thêm mới & cập nhật công thức: Wizard 4 bước (Thông tin chung -> Nguyên liệu -> Các bước -> Ảnh & Xuất bản) tại src/app/dashboard/recipes/new và edit.",
                "- Quản lý form bằng React Hook Form + Zod validation runtime.",
                "- Cơ chế khóa lạc quan OCC qua RowVersion ngăn chặn Lost Update khi 2 tác giả sửa đồng thời.",
                "- Vòng đời công thức: Quản lý tại src/app/dashboard/recipes/page.tsx, chuyển đổi trạng thái Draft -> Publish -> Archived -> Deleted an toàn.",
                "- Quy tắc nghiệp vụ C02: Chỉ cho phép Publish khi công thức có đủ nguyên liệu, bước nấu và ảnh hợp lệ."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend/src/app/dashboard/recipes",
            "progress": "100%"
        },
        {
            "stt": "3",
            "title": "[Danh sách Người dùng] Hoàn thành việc xây dựng trang xem danh sách người dùng (admin, tác giả)",
            "done": [
                "- Backend: Cài đặt endpoint GET /api/v1/users và GET /api/v1/auth/users qua MediatR GetUsersQuery trả về UserDto.",
                "- Frontend: Xây dựng trang danh bạ người dùng tại src/app/dashboard/users/page.tsx.",
                "- Thống kê tổng quan: Tổng thành viên, số lượng Tác giả công thức, số lượng Quản trị viên (Admin).",
                "- Tìm kiếm theo tên/email, lọc theo vai trò (Admin, Tác giả, Thành viên).",
                "- Thẻ người dùng: Avatar, Họ và tên, email xác thực, Role badge, tiểu sử Bio, ngày tham gia, liên kết xem bài viết.",
                "- Tích hợp lối tắt Người dùng trên Header.tsx (desktop & mobile)."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/frontend/src/app/dashboard/users/page.tsx",
            "progress": "100%"
        },
        {
            "stt": "4",
            "title": "[Nhật ký Hệ thống] Bổ sung tính năng logging để ghi lại nhật ký hệ thống",
            "done": [
                "- Ghi log có cấu trúc bằng Serilog và OpenTelemetry vào Console và máy chủ Seq.",
                "- Ghi nhận chi tiết HTTP request: HTTP status, đường dẫn URL, thời gian xử lý (Elapsed ms), User Agent, IP, UserId.",
                "- Phân phối TraceId và tự sinh X-Correlation-ID cho mỗi request, gắn vào LogContext để truy vết xuyên suốt Frontend -> API -> DB.",
                "- Tự động che giấu thông tin nhạy cảm (Password, Token, Secret) không ghi vào log.",
                "- Giám sát sức khỏe hệ thống: Endpoint /health, /health/live, /health/ready kiểm tra kết nối DB và Storage."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/src/backend/CulinaryBlog.API/Program.cs",
            "progress": "100%"
        },
        {
            "stt": "5",
            "title": "[Hoàn thiện Đặc tả SRS] Hoàn thiện các tính năng khác theo đặc tả trong file SRS",
            "done": [
                "- Đáp ứng đủ 34 Functional Requirements (FRs), 30 Non-Functional Requirements (NFRs) và 10 Ràng buộc kỹ thuật (CONS) của SRS v1.1.1.",
                "- Kiến trúc Clean Architecture 4 tầng kết hợp mẫu CQRS / MediatR và FluentValidation.",
                "- Bảo mật toàn diện: CORS dynamic cho Vercel/localhost, Rate Limiting chống DDoS, chuẩn lỗi RFC 9457 Problem Details.",
                "- Bộ đệm đa tầng: Redis Cache-aside, OutputCache và Next.js Incremental Static Regeneration (ISR).",
                "- Tự động hóa kiểm thử: 448/448 test backend pass 100% (443 unit/integration + 5 concurrency spike), 93/93 test frontend Jest pass 100%.",
                "- Script sao lưu & phục hồi CSDL tự động BCP/DR kiểm tra mã băm SHA-256 (thời gian khôi phục 1.8s, RTO < 2 phút, RPO = 0)."
            ],
            "not_done": [
                "- Không có (đã hoàn thành 100% yêu cầu)."
            ],
            "link": "https://github.com/2312741-sudo/PTUDWNC-2026-Nhom16/tree/main/README.md",
            "progress": "100%"
        }
    ]

    create_lab_doc(
        "07 (Tuần 7)",
        "Từ ngày: 15/10/2026\tđến ngày: 21/10/2026",
        lab7_tasks,
        [
            "docs/evidence/TV1/Lab7_2312741_NguyenThanhTam.docx",
            "docs/evidence/TV1/Lab07_2312741_NguyenThanhTam.docx",
            os.path.expanduser("~/Downloads/Lab7_2312741_NguyenThanhTam.docx"),
            os.path.expanduser("~/Downloads/Lab07_2312741_NguyenThanhTam.docx")
        ]
    )

    create_lab_doc(
        "08 (Tuần 8)",
        "Từ ngày: 22/10/2026\tđến ngày: 28/10/2026",
        lab8_tasks,
        [
            "docs/evidence/TV1/Lab8_2312741_NguyenThanhTam.docx",
            "docs/evidence/TV1/Lab08_2312741_NguyenThanhTam.docx",
            os.path.expanduser("~/Downloads/Lab8_2312741_NguyenThanhTam.docx"),
            os.path.expanduser("~/Downloads/Lab08_2312741_NguyenThanhTam.docx")
        ]
    )

if __name__ == "__main__":
    main()
